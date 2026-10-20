#!/usr/bin/env python3
"""POC B1: source-first five-slide narrated video using Piper and FFmpeg.
Use --tts espeak only for a mechanical fallback preview (NOT a Piper benchmark).
"""
import argparse, hashlib, json, pathlib, shutil, subprocess, sys, wave
from pptx import Presentation

ROOT=pathlib.Path(__file__).resolve().parent
DECK=ROOT/'source'/'codex-agent-skills-and-adapters-detailed.pptx'

def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)

def extract(deck=DECK):
    p=Presentation(deck)
    assert len(p.slides)>=13
    for idx in range(9,14):
        s=p.slides[idx-1]
        imgs=[shape for shape in s.shapes if shape.shape_type==13]
        if len(imgs)!=1: raise RuntimeError(f'Slide {idx}: expected 1 image, got {len(imgs)}')
        (ROOT/'slides'/f'{idx:02d}.png').write_bytes(imgs[0].image.blob)
        n=s.notes_slide.notes_text_frame.text.strip()
        (ROOT/'notes'/f'{idx:02d}.txt').write_text(n+'\n',encoding='utf8')


def probe(path):
    q=subprocess.check_output(['ffprobe','-v','error','-show_entries','format=duration','-of','default=nokey=1:noprint_wrappers=1',str(path)],text=True).strip()
    return float(q)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--tts',choices=['piper','espeak'],default='piper')
    ap.add_argument('--piper-model',help='Piper .onnx voice model path, required for --tts piper')
    ap.add_argument('--extract-only',action='store_true')
    ap.add_argument('--deck',type=pathlib.Path,default=DECK,help='Source detailed PPTX containing slides 9-13')
    args=ap.parse_args()
    for p in ('slides','notes','build'):(ROOT/p).mkdir(exist_ok=True)
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):sys.exit('FFmpeg and FFprobe are required')
    if not args.deck.is_file(): sys.exit(f"Source PPTX not found: {args.deck}")
    extract(args.deck)
    if args.extract_only:
        print('Extracted last five visuals and verbatim PowerPoint speaker notes.')
        return
    if args.tts=='piper':
        if not shutil.which('piper'):sys.exit('Piper command not installed. Install piper-tts and retry.')
        if not args.piper_model or not pathlib.Path(args.piper_model).is_file():sys.exit('Provide --piper-model PATH to a downloaded Piper .onnx model')
    else:
        if not shutil.which('espeak'):sys.exit('eSpeak command not installed')
    meta=[]
    for idx in range(9,14):
        note=(ROOT/'notes'/f'{idx:02d}.txt').read_text(encoding='utf8').strip()
        wav=ROOT/'build'/f'{idx:02d}.wav'
        if args.tts=='piper':
            run(['piper','--model',args.piper_model,'--output_file',str(wav)],input=(note+'\n').encode('utf8'))
        else:
            run(['espeak','-v','en-us','-s','155','-f',str(ROOT/'notes'/f'{idx:02d}.txt'),'-w',str(wav)])
        duration=probe(wav)+0.65
        part=ROOT/'build'/f'{idx:02d}.mp4'
        run(['ffmpeg','-hide_banner','-loglevel','error','-y','-loop','1','-framerate','1','-i',str(ROOT/'slides'/f'{idx:02d}.png'),'-i',str(wav),'-vf','scale=1280:720:force_original_aspect_ratio=decrease,pad=1280:720:(ow-iw)/2:(oh-ih)/2,format=yuv420p','-c:v','libx264','-preset','ultrafast','-crf','26','-r','12','-c:a','aac','-b:a','128k','-t',str(duration),'-af','apad','-movflags','+faststart',str(part)])
        meta.append({'slide':idx,'notes_sha256':hashlib.sha256(note.encode()).hexdigest(),'characters':len(note),'audio_seconds':round(probe(wav),2),'video_seconds':round(probe(part),2)})
    concat=ROOT/'build'/'concat.txt'
    concat.write_text(''.join("file '%s'\n" % (ROOT/'build'/f'{idx:02d}.mp4') for idx in range(9,14)))
    output=ROOT/'build'/f'last-five-slides-{args.tts}.mp4'
    run(['ffmpeg','-hide_banner','-loglevel','error','-y','-f','concat','-safe','0','-i',str(concat),'-c','copy','-movflags','+faststart',str(output)])
    result={'engine':args.tts,'disclaimer':'eSpeak is mechanical fallback, not Piper voice assessment' if args.tts=='espeak' else 'Piper generated voice. Requires human listening acceptance.','slides':meta,'output':str(output),'duration_seconds':round(probe(output),2)}
    (ROOT/'build'/f'manifest-{args.tts}.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))
if __name__=='__main__':main()
