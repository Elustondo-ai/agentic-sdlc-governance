"""Deterministic PPTX export from validated sources; no source modifications."""
import hashlib
import io
import json
import os
import platform
import tempfile
from datetime import datetime
from importlib.metadata import version
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import cairosvg
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches
from contracts import BuildResult, PresentationError, SlideSource
from sources import discover


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def notes_text(slide: SlideSource) -> str:
    text = slide.narration
    if slide.references:
        text += "\n\n--- Reference metadata (not narration) ---\n" + slide.references
    return text


def stable_pptx(prs: Presentation) -> bytes:
    prs.core_properties.created = prs.core_properties.modified = datetime(2000, 1, 1)
    prs.core_properties.last_modified_by = "ASDG presentation builder v0.2"
    raw, output = io.BytesIO(), io.BytesIO()
    prs.save(raw)
    with ZipFile(raw) as source, ZipFile(output, "w", ZIP_DEFLATED, compresslevel=6) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return output.getvalue()


def atomic_write(path: Path, data: bytes) -> None:
    if path.is_symlink():
        raise PresentationError(f"Refusing symlink output: {path}")
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=".asdg-")
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def build(presentation: Path, profile: str = "intro") -> BuildResult:
    slides = discover(presentation, profile)
    presentation = presentation.resolve()
    output = presentation / "build"
    if output.is_symlink():
        raise PresentationError(f"Build directory cannot be a symlink: {output}")
    if output.exists() and not output.is_dir():
        raise PresentationError(f"Build output is not a directory: {output}")
    original = {p: digest(p) for s in slides for p in (s.visual, *(s.note_files or (s.notes,)))}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = presentation.name
    prs.core_properties.subject = "Derived from versioned SVG and Markdown sources"
    previews, warnings, entries = [], [], []
    for source in slides:
        png = cairosvg.svg2png(bytestring=source.visual.read_bytes(), output_width=2400)
        im = Image.open(io.BytesIO(png)).convert("RGB")
        width, height = im.size
        scale = min(prs.slide_width / width, prs.slide_height / height)
        w, h = round(width * scale), round(height * scale)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_picture(io.BytesIO(png),
            (prs.slide_width - w) // 2, (prs.slide_height - h) // 2, width=w, height=h)
        shape.name = source.title
        slide.notes_slide.notes_text_frame.text = notes_text(source)
        previews.append(ImageOps.pad(im, (480, 270), color="white"))
        if source.raster_backed:
            warnings.append(f"{source.folder.name}: raster-backed SVG preserved; not fully editable")
        entries.append({"number": source.number, "folder": source.folder.name,
            "title": source.title, "visual": source.visual.name, "notes": source.notes.name,
            "visual_sha256": original[source.visual], "notes_sha256": original[source.notes],
            "raster_backed": source.raster_backed,
            "note_files": [{"file": p.name, "sha256": original[p]}
                           for p in (source.note_files or (source.notes,))],
            "narration_words": len(source.narration.split())})
    if any(digest(p) != value for p, value in original.items()):
        raise PresentationError("Source changed while building; no output was published")
    pptx_bytes = stable_pptx(prs)
    reread = Presentation(io.BytesIO(pptx_bytes))
    if len(reread.slides) != len(slides) or any(
        a.notes_slide.notes_text_frame.text != notes_text(b) for a, b in zip(reread.slides, slides)
    ):
        raise PresentationError("PPTX round-trip note/count verification failed")
    output.mkdir(exist_ok=True)
    suffix = "-detailed" if profile == "detailed" else ""
    target = output / f"{presentation.name}{suffix}.pptx"
    manifest = output / f"build-manifest{suffix}.json"
    report = {"builder_version": "0.2.0", "narration_profile": profile, "presentation": presentation.name,
        "slide_count": len(slides), "output": target.name,
        "pptx_sha256": hashlib.sha256(pptx_bytes).hexdigest(),
        "dependencies": {p: version(p) for p in ("python-pptx", "CairoSVG", "Pillow", "defusedxml")},
        "python": platform.python_version(), "slides": entries, "warnings": warnings,
        "notes_round_trip_verified": True, "source_files_unchanged": True}
    atomic_write(target, pptx_bytes)
    columns = min(3, len(previews)); rows = (len(previews) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * 500, rows * 305), "#eeeeee")
    draw = ImageDraw.Draw(sheet)
    for index, image in enumerate(previews):
        x, y = (index % columns) * 500 + 10, (index // columns) * 305 + 10
        sheet.paste(image, (x, y)); draw.text((x, y + 275), f"{index+1:02d}", fill="black")
    thumb = io.BytesIO(); sheet.save(thumb, "PNG")
    atomic_write(output / f"preview{suffix}.png", thumb.getvalue())
    # A read-only assembly for narration tools; authoritative parts stay by each slide.
    script = "\n\n".join(f"# {s.title}\n\n{s.narration}" for s in slides) + "\n"
    atomic_write(output / f"{presentation.name}{suffix}-notes.md", script.encode())
    atomic_write(manifest, (json.dumps(report, indent=2, ensure_ascii=False) + "\n").encode())
    return BuildResult(presentation, target, manifest, len(slides), tuple(warnings))
