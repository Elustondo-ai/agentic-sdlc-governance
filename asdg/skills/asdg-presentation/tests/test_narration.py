"""Tests for ordered narration, explicit profiles, and preserved intro output."""
import json
from pathlib import Path
from pptx import Presentation
import unittest
import test_build
MD = test_build.MD
from contracts import PresentationError
from sources import discover
from export import build


class LayeredNarrationTests(unittest.TestCase):
    setUp = test_build.PresentationTests.setUp
    def add_parts(self, details=True):
        (self.slide/'speaker-notes-02-walkthrough.md').write_text('# Walkthrough\n\nRead the visual.\n')
        if details:
            (self.slide/'speaker-notes-03-details.md').write_text('# Details\n\nExplain the example.\n')

    def test_intro_ignores_optional_parts(self):
        before = build(self.deck).pptx.read_bytes()
        self.add_parts()
        self.assertEqual(before, build(self.deck).pptx.read_bytes())

    def test_detailed_note_order(self):
        self.add_parts()
        result = build(self.deck, 'detailed')
        text = Presentation(result.pptx).slides[0].notes_slide.notes_text_frame.text
        self.assertEqual(text, 'First paragraph.\n\nSecond paragraph.\n\nRead the visual.\n\nExplain the example.')
        self.assertEqual(result.pptx.name, 'sample-deck-detailed.pptx')

    def test_two_notes_supported(self):
        self.add_parts(details=False)
        slides = discover(self.deck, 'detailed')
        self.assertEqual(len(slides[0].note_files), 2)

    def test_missing_walkthrough_rejected(self):
        (self.slide/'speaker-notes-03-details.md').write_text(MD)
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_detailed_needs_walkthrough(self):
        with self.assertRaises(PresentationError): build(self.deck, 'detailed')
        self.assertFalse((self.deck/'build').exists())

    def test_detailed_preserves_existing_intro_outputs(self):
        first = build(self.deck)
        baseline = {p.name:p.read_bytes() for p in first.pptx.parent.iterdir()}
        self.add_parts()
        build(self.deck, 'detailed')
        for name, value in baseline.items():
            self.assertEqual(value, (first.pptx.parent/name).read_bytes())

    def test_manifest_hashes_all_selected_parts(self):
        self.add_parts()
        result = build(self.deck, 'detailed')
        report = json.loads(result.manifest.read_text())
        self.assertEqual(report['narration_profile'], 'detailed')
        self.assertEqual([x['file'] for x in report['slides'][0]['note_files']],
            ['speaker-notes.md','speaker-notes-02-walkthrough.md','speaker-notes-03-details.md'])

    def test_unknown_profile_rejected(self):
        with self.assertRaises(PresentationError): discover(self.deck, 'invented')

    def test_blank_supplement_rejected(self):
        (self.slide/'speaker-notes-02-walkthrough.md').write_text('# Heading only\n')
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_unrecognized_note_file_rejected(self):
        (self.slide/'some-more-notes.md').write_text(MD)
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_new_intro_name_supported(self):
        (self.slide/'speaker-notes.md').rename(self.slide/'speaker-notes-01-intro.md')
        self.add_parts()
        self.assertEqual(len(discover(self.deck, 'detailed')[0].note_files), 3)

    def test_supplement_metadata_not_spoken(self):
        self.add_parts(details=False)
        p = self.slide/'speaker-notes-02-walkthrough.md'
        p.write_text(p.read_text()+'\n<!-- private reference -->\n')
        slide = discover(self.deck,'detailed')[0]
        self.assertNotIn('private reference',slide.narration)
        self.assertIn('private reference',slide.references)

    def test_detailed_repeat_build(self):
        self.add_parts()
        before = build(self.deck, 'detailed').pptx.read_bytes()
        self.assertEqual(before, build(self.deck,'detailed').pptx.read_bytes())
