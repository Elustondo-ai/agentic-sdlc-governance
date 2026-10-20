"""Isolated tests for presentation source contracts and derived output."""
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(SKILL / 'types'), str(SKILL / 'src')]
from contracts import PresentationError
from sources import discover, read_notes, safe_svg
from export import build

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="160" height="90"><rect width="160" height="90" fill="white"/><text x="10" y="40">Example</text></svg>'
MD = 'AI-MODIFIED - 2026-10-08T21:03Z\n\n# Example\n\nFirst paragraph.\n\nSecond paragraph.\n'

class PresentationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.deck = Path(self.temp.name) / 'sample-deck'
        self.slide = self.deck / '01-example'
        self.slide.mkdir(parents=True)
        (self.slide/'slide.svg').write_text(SVG)
        (self.slide/'speaker-notes.md').write_text(MD)

    def test_valid_pair(self):
        self.assertEqual(len(discover(self.deck)), 1)

    def test_original_sources_unchanged(self):
        before = {p.name: p.read_bytes() for p in self.slide.iterdir()}
        build(self.deck)
        self.assertEqual(before, {p.name: p.read_bytes() for p in self.slide.iterdir()})

    def test_round_trip_notes(self):
        result = build(self.deck)
        deck = Presentation(result.pptx)
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(deck.slides[0].notes_slide.notes_text_frame.text,
                         'First paragraph.\n\nSecond paragraph.')
        self.assertEqual(result.pptx.parent, self.deck/'build')

    def test_repeat_build_same_bytes(self):
        first = build(self.deck).pptx.read_bytes()
        self.assertEqual(first, build(self.deck).pptx.read_bytes())

    def test_missing_notes(self):
        (self.slide/'speaker-notes.md').unlink()
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_extra_visual(self):
        (self.slide/'extra.svg').write_text(SVG)
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_duplicate_notes(self):
        (self.slide/'speaker-text.md').write_text(MD)
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_missing_number(self):
        self.slide.rename(self.deck/'02-example')
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_unexpected_subfolder(self):
        (self.deck/'unlabelled').mkdir()
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_unsupported_renderer(self):
        (self.slide/'slide.svg').rename(self.slide/'slide.mmd')
        with self.assertRaisesRegex(PresentationError, 'SVG only'): discover(self.deck)

    def test_external_svg_reference(self):
        (self.slide/'slide.svg').write_text(SVG.replace('</svg>', '<image href="https://example.invalid/private.png"/></svg>'))
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_external_css_resource(self):
        (self.slide/'slide.svg').write_text(SVG.replace('</svg>', '<style>text{fill:url(https://example.invalid)}</style></svg>'))
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_svg_script(self):
        (self.slide/'slide.svg').write_text(SVG.replace('</svg>', '<script>bad()</script></svg>'))
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_empty_narration(self):
        (self.slide/'speaker-notes.md').write_text('# Empty\n')
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_reference_metadata_separate(self):
        notes = self.slide/'speaker-notes.md'
        notes.write_text(MD + '\n<!-- source URL, not narration -->\n## Fine print\nSource attribution.\n')
        _, narration, references = read_notes(notes)
        self.assertNotIn('Source attribution', narration)
        self.assertIn('Source attribution', references)
        self.assertIn('source URL', references)

    def test_build_symlink_rejected(self):
        elsewhere = Path(self.temp.name)/'elsewhere'; elsewhere.mkdir()
        (self.deck/'build').symlink_to(elsewhere, target_is_directory=True)
        with self.assertRaises(PresentationError): build(self.deck)
        self.assertEqual(list(elsewhere.iterdir()), [])

    def test_source_symlink_rejected(self):
        notes = self.slide/'speaker-notes.md'
        target = Path(self.temp.name)/'notes.md'; notes.rename(target); notes.symlink_to(target)
        with self.assertRaises(PresentationError): discover(self.deck)

    def test_no_outputs_on_invalid_source(self):
        (self.slide/'speaker-notes.md').unlink()
        with self.assertRaises(PresentationError): build(self.deck)
        self.assertFalse((self.deck/'build').exists())

if __name__ == '__main__': unittest.main()
